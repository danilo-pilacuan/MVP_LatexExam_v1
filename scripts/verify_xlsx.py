"""Verificación del Excel exportado (corre dentro del contenedor)."""
import sys
sys.path.insert(0, "/app")
from openpyxl import load_workbook

wb = load_workbook("/app/output/preguntas_20260917_071733.xlsx")
ws = wb.active
print("Filas:", ws.max_row, "Columnas:", ws.max_column)

# Buscar una fila con opciones (opcion_multiple)
found = False
for row in ws.iter_rows(min_row=2, values_only=True):
    if row[6] == "opcion_multiple":
        print("Pregunta:", str(row[7])[:60])
        print("Opciones:", row[8])
        print("Respuesta:", row[9])
        print("Verif humano:", row[12])
        found = True
        break
if not found:
    print("No hay filas con opcion_multiple")

# Contar tipos de pregunta
from collections import Counter
tipos = Counter(row[6] for row in ws.iter_rows(min_row=2, values_only=True))
print("Tipos:", dict(tipos))
